"""Unit tests for the GENERAL source-backed field drift core (#6).

Exercises the content-control read/write helpers directly (no DB): build a DOCX with
tagged content controls exactly like replace_placeholders_in_docx produces, then verify
read_filled_field_values reads them back and set_filled_field_values updates by tag.
"""
import tempfile
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from app.modules.agreements.services.placeholder_fill import (  # noqa: E402
    read_filled_field_values,
    set_filled_field_values,
)


def _add_content_control(paragraph, tag: str, value: str) -> None:
    """Append a tagged content control (w:sdt) to a paragraph — mirrors the structure
    _replace_with_editable_content_control writes (w:sdtPr/w:tag + w:sdtContent/w:r/w:t)."""
    p_elem = paragraph._p
    sdt = OxmlElement("w:sdt")
    sdt_pr = OxmlElement("w:sdtPr")
    tag_el = OxmlElement("w:tag")
    tag_el.set(qn("w:val"), tag)
    sdt_pr.append(tag_el)
    sdt_content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = value
    run.append(t)
    sdt_content.append(run)
    sdt.append(sdt_pr)
    sdt.append(sdt_content)
    p_elem.append(sdt)


def _make_docx(fields: dict) -> str:
    doc = Document()
    for tag, value in fields.items():
        p = doc.add_paragraph()
        _add_content_control(p, tag, value)
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    tmp.close()
    doc.save(tmp.name)
    return tmp.name


def test_read_filled_field_values_reads_tags_uppercased():
    path = _make_docx({"hospital_name": "Old General", "PI_NAME": "Dr. A"})
    values = read_filled_field_values(path)
    # keyed UPPER-CASE regardless of the tag's original case
    assert values.get("HOSPITAL_NAME") == "Old General"
    assert values.get("PI_NAME") == "Dr. A"
    Path(path).unlink(missing_ok=True)


def test_set_filled_field_values_updates_only_targeted_tags():
    path = _make_docx({"HOSPITAL_NAME": "Old General", "PI_NAME": "Dr. A"})
    out = set_filled_field_values(path, {"hospital_name": "New General Hospital"})
    values = read_filled_field_values(out)
    assert values.get("HOSPITAL_NAME") == "New General Hospital"  # pulled from source
    assert values.get("PI_NAME") == "Dr. A"                       # untouched edit preserved
    Path(path).unlink(missing_ok=True)
    Path(out).unlink(missing_ok=True)


def test_read_filled_field_values_missing_file_is_empty():
    assert read_filled_field_values("/no/such/file.docx") == {}
